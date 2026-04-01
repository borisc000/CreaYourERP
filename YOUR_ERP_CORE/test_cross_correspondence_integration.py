import asyncio
import base64
from io import BytesIO
import unittest

from docx import Document

from core.YOUR_ERP_core_framework import CoreFramework, Request
from core.YOUR_ERP_orm import BaseModel
from modules.base.module_base import BaseModule as BaseAppModule, Company, User
from modules.crm.module_crm import CRMModule, Customer, Lead, ServiceType
from modules.document_center.module_document_center import (
    DocumentCenterModule,
    DocumentTemplate,
    GeneratedDocument,
)
from modules.hr.module_hr import (
    AccreditationRequirement,
    EmployeeAccreditationDocument,
    EmployeeContract,
    EmployeeProfile,
    HRModule,
)
from modules.safety.module_safety import SafetyFolder, SafetyModule
from modules.signature.module_signature import SignatureModule


def _docx_template_base64(*lines: str) -> str:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    buffer = BytesIO()
    doc.save(buffer)
    return base64.b64encode(buffer.getvalue()).decode("utf-8")


class CrossCorrespondenceIntegrationTest(unittest.TestCase):
    maxDiff = None

    def setUp(self) -> None:
        BaseModel._store = {}
        BaseModel._id_counters = {}

        self.framework = CoreFramework(
            {
                "database_url": "memory://demo",
                "debug": False,
                "secret_key": "test-secret",
                "modules_to_load": [
                    "base",
                    "signature",
                    "crm",
                    "hr",
                    "safety",
                    "document_center",
                ],
            }
        )
        for module_class in (
            BaseAppModule,
            SignatureModule,
            CRMModule,
            HRModule,
            SafetyModule,
            DocumentCenterModule,
        ):
            self.framework.register_module_class(module_class)
        self.framework.initialize()

        self.company = Company.create(
            {
                "name": "Constructora Demo",
                "email": "demo@example.com",
                "legal_name": "Constructora Demo SpA",
                "tax_id": "76.123.456-7",
            }
        )

        self.user = User(
            company_id=self.company.id,
            email="admin@example.com",
            name="Admin Demo",
            role="company_admin",
            is_admin=True,
            allowed_modules=["document_center", "safety", "hr", "crm", "signature"],
        )
        self.user.set_password("secret123")
        self.user.save()

        self.customer = Customer.create(
            {
                "name": "Cliente Uno",
                "email": "cliente@example.com",
                "company_id": self.company.id,
            }
        )
        self.service_type = ServiceType.create(
            {
                "name": "Montaje industrial",
                "company_id": self.company.id,
            }
        )
        self.lead = Lead.create(
            {
                "project_code": "PRJ-5001",
                "title": "Parada de planta",
                "customer_id": self.customer.id,
                "service_type_id": self.service_type.id,
                "company_id": self.company.id,
                "priority": "medium",
                "status": "open",
                "probability": 60,
            }
        )
        self.employee = EmployeeProfile.create(
            {
                "full_name": "Juan Perez",
                "company_id": self.company.id,
                "status": "active",
                "position_title": "Prevencionista",
                "work_email": "juan.perez@example.com",
            }
        )
        self.contract = EmployeeContract.create(
            {
                "employee_id": self.employee.id,
                "company_id": self.company.id,
                "contract_type": "indefinite",
                "status": "active",
                "start_date": "2026-03-01",
                "salary_amount": 950000,
            }
        )
        self.folder = SafetyFolder.create(
            {
                "lead_id": self.lead.id,
                "company_id": self.company.id,
                "status": "in_progress",
                "assigned_employee_ids": [self.employee.id],
            }
        )

    def _dispatch(self, method: str, path: str, data=None, params=None):
        request = Request(
            path=path,
            method=method,
            user_id=self.user.id,
            company_id=self.company.id,
            data=data or {},
            params=params or {},
        )
        return asyncio.run(self.framework.dispatch_request(request))

    def _create_template(
        self,
        *,
        name: str,
        target_module: str,
        document_type: str,
        category: str,
        placeholders,
        accreditation_code: str,
        accreditation_category: str,
    ) -> DocumentTemplate:
        template_data = _docx_template_base64(
            f"{name}",
            "Trabajador: <<nombre>>",
            "Cliente: <<cliente>>",
            "Contexto: <<source_label>>",
            "Items: <<detail_items_multiline>>",
        )
        return DocumentTemplate.create(
            {
                "name": name,
                "description": f"Plantilla {name}",
                "category": category,
                "document_type": document_type,
                "target_module": target_module,
                "status": "active",
                "company_id": self.company.id,
                "requires_signature": False,
                "auto_register_accreditation": True,
                "accreditation_requirement_code": accreditation_code,
                "accreditation_category": accreditation_category,
                "filename_pattern": f"{name} <<nombre>>",
                "original_filename": f"{name.lower().replace(' ', '_')}.docx",
                "template_data": template_data,
                "placeholder_keys": list(placeholders),
                "preview_text": name,
                "tags": [category],
            }
        )

    def test_worker_generation_registers_accreditation_and_close_syncs_status(self):
        template = self._create_template(
            name="Anexo indefinido",
            target_module="hr",
            document_type="anexo indefinido",
            category="rrhh",
            placeholders=["nombre", "cliente", "source_label"],
            accreditation_code="ANEXO_INDEFINIDO",
            accreditation_category="contractual",
        )

        response = self._dispatch(
            "POST",
            "/document-center/worker-generate",
            data={
                "employee_id": self.employee.id,
                "template_ids": [template.id],
                "customer_id": self.customer.id,
                "lead_id": self.lead.id,
                "source_module": "hr",
                "source_record_id": self.employee.id,
                "source_label": "Ficha de trabajador",
                "target_module": "hr",
                "target_record_id": self.employee.id,
                "detail_items": ["Anexo indefinido"],
                "document_date": "2026-03-30",
            },
        )

        self.assertEqual(response.status, 201, response.errors)
        self.assertEqual(response.data["summary"]["documents_generated"], 1)
        self.assertEqual(response.data["summary"]["accreditation_registered"], 1)

        generated_payload = response.data["generated_documents"][0]
        self.assertEqual(generated_payload["source_module"], "hr")
        self.assertEqual(generated_payload["target_module"], "hr")
        self.assertEqual(generated_payload["target_record_id"], self.employee.id)
        self.assertTrue(generated_payload["accreditation_document_id"])
        self.assertEqual(
            generated_payload["workspace_url"],
            f"/app/cross-correspondence?generated_document_id={generated_payload['id']}",
        )

        generated = GeneratedDocument.find_by_id(generated_payload["id"])
        self.assertIsNotNone(generated)
        self.assertEqual(generated.employee_id, self.employee.id)
        self.assertEqual(generated.customer_id, self.customer.id)

        requirements = AccreditationRequirement.search([("company_id", "=", self.company.id)])
        self.assertEqual(len(requirements), 1)
        self.assertEqual(requirements[0].code, "ANEXO_INDEFINIDO")
        self.assertEqual(requirements[0].customer_id, self.customer.id)

        accreditation_doc = EmployeeAccreditationDocument.find_by_id(
            generated.accreditation_document_id
        )
        self.assertIsNotNone(accreditation_doc)
        self.assertEqual(accreditation_doc.source_module, "hr")
        self.assertEqual(accreditation_doc.verification_status, "pending_review")
        self.assertEqual(
            accreditation_doc.document_url,
            f"/app/cross-correspondence?generated_document_id={generated.id}",
        )

        approve_response = self._dispatch(
            "POST", f"/document-center/generated/{generated.id}/approve", data={}
        )
        self.assertEqual(approve_response.status, 200, approve_response.errors)

        close_response = self._dispatch(
            "POST", f"/document-center/generated/{generated.id}/close", data={}
        )
        self.assertEqual(close_response.status, 200, close_response.errors)

        closed_generated = GeneratedDocument.find_by_id(generated.id)
        self.assertEqual(closed_generated.status, "closed")

        synced_accreditation = EmployeeAccreditationDocument.find_by_id(accreditation_doc.id)
        self.assertEqual(synced_accreditation.verification_status, "approved")
        self.assertEqual(synced_accreditation.verified_by, self.user.id)
        self.assertTrue(synced_accreditation.verified_at)

    def test_safety_ppe_delivery_calls_document_center_and_links_results(self):
        template = self._create_template(
            name="Entrega EPP",
            target_module="safety",
            document_type="entrega epp",
            category="safety",
            placeholders=["nombre", "cliente", "detail_items_multiline", "source_label"],
            accreditation_code="ENTREGA_EPP",
            accreditation_category="safety",
        )

        response = self._dispatch(
            "POST",
            f"/safety/folders/{self.folder.id}/ppe-deliveries",
            data={
                "employee_id": self.employee.id,
                "delivery_date": "2026-03-30",
                "status": "delivered",
                "items": ["Casco", "Guantes", "Lentes"],
                "notes": "Entrega inicial de faena",
                "document_template_id": template.id,
            },
        )

        self.assertEqual(response.status, 201, response.errors)
        self.assertFalse(response.data["document_generation_error"])
        self.assertEqual(len(response.data["generated_documents"]), 1)

        delivery_id = response.data["delivery"]["id"]
        generated_payload = response.data["generated_documents"][0]
        self.assertEqual(generated_payload["source_module"], "safety")
        self.assertEqual(generated_payload["source_record_id"], delivery_id)
        self.assertEqual(generated_payload["target_module"], "safety")
        self.assertEqual(generated_payload["target_record_id"], self.folder.id)
        self.assertTrue(generated_payload["accreditation_document_id"])

        list_response = self._dispatch(
            "GET",
            "/document-center/generated",
            params={
                "target_module": "safety",
                "target_record_id": str(self.folder.id),
                "source_module": "safety",
                "source_record_id": str(delivery_id),
            },
        )
        self.assertEqual(list_response.status, 200, list_response.errors)
        self.assertEqual(list_response.data["count"], 1)
        self.assertEqual(list_response.data["results"][0]["id"], generated_payload["id"])

        accreditation_doc = EmployeeAccreditationDocument.find_by_id(
            generated_payload["accreditation_document_id"]
        )
        self.assertIsNotNone(accreditation_doc)
        self.assertEqual(accreditation_doc.source_module, "safety")
        self.assertEqual(accreditation_doc.employee_id, self.employee.id)


if __name__ == "__main__":
    unittest.main()
